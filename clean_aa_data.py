"""Clean the combined Assessment Agency master against the NCVET master lists.

Stage 3 of the monthly pipeline:

    Google link sheet -> aa_monitor.py -> clean_aa_data.py -> monthly_scoring_report.py

The combiner (aa_monitor.py / combine_aa_files.py) aligns every submitted row to
the 41 DCF columns but keeps the values exactly as the Assessment Agency typed
them. This script standardises those values against the approved master lists,
normalises every date and number, and writes a full audit trail of what it
changed.

Edit the USER SETTINGS block below, then run:
    python clean_aa_data.py

Command-line overrides are also available:
    python clean_aa_data.py --help
"""

from __future__ import annotations

# =============================================================================
# USER SETTINGS - edit these paths/options before running the script
# =============================================================================
from pathlib import Path

# The combined workbook produced by aa_monitor.py or combine_aa_files.py.
INPUT_WORKBOOK = Path(
    r"C:\Users\Anurag\Desktop\EY\Engagements\NCVET\Monitoring\AA_Consolidated_Master.xlsx"
)

# Blank = auto-detect "Master Data" / "Data" / the largest sheet.
INPUT_SHEET_NAME = ""

# Where the cleaned workbook is written. Feed this file to monthly_scoring_report.py.
OUTPUT_FILE = Path(
    r"C:\Users\Anurag\Desktop\EY\Engagements\NCVET\Monitoring\AA_Cleaned_Master.xlsx"
)

# ---- Master reference lists -------------------------------------------------
MONITORING_FOLDER = Path(__file__).resolve().parent

AA_LIST_PATH = MONITORING_FOLDER / "AA List.xlsx"
SECTORS_LIST_PATH = MONITORING_FOLDER / "Sectors_List.xlsx"
STATE_DISTRICT_PATH = MONITORING_FOLDER / "State District List.xlsx"

# The DCF supplies the controlled vocabularies on its LOOKUP sheet (funding
# type, training type, mode of assessment, language, level and so on) and the
# state/district list on its LGD sheet. Leave blank to skip it.
DCF_TEMPLATE_PATH = MONITORING_FOLDER / "AA Monthly DCF 2627 - V1 - 06 May 2026  (1).xlsx"

# ---- Cleaning behaviour -----------------------------------------------------
# Autofix mode: every value that can be resolved to a master-list entry is
# rewritten, and every change is recorded in the Cleaning Log sheet.
FUZZY_CUTOFF = 0.72           # general similarity floor for a rewrite
FUZZY_CUTOFF_AGENCY = 0.80    # agency names are matched more strictly
FUZZY_CUTOFF_DISTRICT = 0.78  # districts are matched inside their state first

# Explicit user-confirmed names take precedence over fuzzy matching. Alias
# targets may differ from the approved master; scoring reconciliation exposes this.
AGENCY_NAME_ALIASES = {
    "FFCPL": "Fashion Future First Assessment Agency",
    "Fashion Future First Assessment Agency": "Fashion Future First Assessment Agency",
}

# Dates in the DCF are day-first (DD-MM-YYYY). Set False only if the agencies
# start submitting US-style month-first dates.
PREFER_DAY_FIRST_DATES = True

# Fills a blank Month/Year of Batch Assessment. Blank = use the reporting month
# recorded by the combiner for that file, then leave blank.
REPORTING_MONTH_HINT = ""     # e.g. "July 2026" or "Jul-26"

# The Cleaning Log can get long. The Change Summary sheet always shows the full
# totals even when the row-level log is truncated.
MAX_LOG_ROWS = 50000

# Keep the pre-cleaning value of these fields as extra columns on Master Data.
KEEP_ORIGINAL_COLUMNS = True
# =============================================================================

import argparse
import json
import logging
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from typing import Any, Iterable, Sequence

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "Missing dependency 'openpyxl'. Run: pip install -r requirements.txt"
    ) from exc


APP_NAME = "NCVET AA Data Cleaner"
APP_VERSION = "1.0.0"

NAVY = "132B46"
LIGHT_BLUE = "D9EAF7"
GREEN = "C6E0B4"
YELLOW = "FFF2CC"
RED = "F4CCCC"
WHITE = "FFFFFF"

MONTH_NAMES = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12,
}
MONTH_LABELS = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]


# -----------------------------------------------------------------------------
# Text helpers
# -----------------------------------------------------------------------------
def clean_text(value: Any) -> str:
    """Trim, collapse whitespace and remove non-breaking/control characters."""
    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN
        return ""
    text = unicodedata.normalize("NFKC", str(value)).replace("\xa0", " ")
    text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or ord(ch) >= 32)
    return " ".join(text.split()).strip()


def normalize(value: Any) -> str:
    """Lower-case, ASCII-fold, de-punctuate. Used for all master-list matching."""
    text = clean_text(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = text.replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", normalize(value))


def display_value(value: Any) -> str:
    """A stable printable form, so a type change alone is not logged as a change."""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and float(value).is_integer():
        return str(int(value))
    return clean_text(value)


AGENCY_WORD_MAP = {
    "private": "pvt", "limited": "ltd", "technologies": "technology",
    "services": "service", "solutions": "solution", "consultants": "consultant",
    "consultancy": "consultant", "skilling": "skill", "skills": "skill",
    "assessments": "assessment", "centres": "centre", "center": "centre",
    "centers": "centre", "india": "india", "incorporated": "inc",
    "corporation": "corp", "company": "co", "organisation": "organization",
}
AGENCY_STOPWORDS = {"the", "a", "an", "of", "and"}


def normalize_agency(value: Any) -> str:
    text = normalize(value)
    text = re.sub(r"\(.*?\)", " ", text)  # drop parenthetical qualifiers
    words = [AGENCY_WORD_MAP.get(word, word) for word in text.split()]
    return " ".join(word for word in words if word not in AGENCY_STOPWORDS).strip()


def similarity(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    base = SequenceMatcher(None, left, right).ratio()
    left_tokens, right_tokens = set(left.split()), set(right.split())
    if left_tokens and right_tokens:
        overlap = len(left_tokens & right_tokens) / len(left_tokens | right_tokens)
        base = max(base, (base + overlap) / 2)
    if left in right or right in left:
        shorter, longer = sorted((left, right), key=len)
        # A short substring ("una" inside "puna") is not evidence of a match.
        if len(shorter) >= 5 and len(shorter) / len(longer) >= 0.6:
            base = max(base, 0.90)
    return base


# -----------------------------------------------------------------------------
# Vocabulary matcher
# -----------------------------------------------------------------------------
@dataclass
class Match:
    value: Any
    method: str
    score: float


class Vocabulary:
    """A canonical list of allowed values with exact, alias and fuzzy matching."""

    def __init__(
        self,
        name: str,
        values: Iterable[str],
        aliases: dict[str, str] | None = None,
        cutoff: float = FUZZY_CUTOFF,
        normalizer=normalize,
    ) -> None:
        self.name = name
        self.normalizer = normalizer
        self.cutoff = cutoff
        self.values: list[str] = []
        self.by_normalized: dict[str, str] = {}
        for item in values:
            text = clean_text(item)
            key = self.normalizer(text)
            if not text or not key or key in self.by_normalized:
                continue
            self.values.append(text)
            self.by_normalized[key] = text
        self.aliases: dict[str, str] = {}
        for alias, target in (aliases or {}).items():
            alias_key = self.normalizer(alias)
            canonical = self.by_normalized.get(self.normalizer(target), clean_text(target))
            if alias_key and canonical:
                self.aliases[alias_key] = canonical
        self._cache: dict[str, Match] = {}

    def __bool__(self) -> bool:
        return bool(self.values)

    def best_match(self, raw: Any) -> Match:
        """The closest candidate regardless of the cutoff."""
        text = clean_text(raw)
        if not text:
            return Match("", "Blank", 0.0)
        key = self.normalizer(text)
        if key in self.by_normalized:
            canonical = self.by_normalized[key]
            return Match(canonical, "Exact" if canonical == text else "Normalized", 1.0)
        if key in self.aliases:
            return Match(self.aliases[key], "Alias", 1.0)
        best_value, best_score = "", 0.0
        for candidate_key, candidate in self.by_normalized.items():
            score = similarity(key, candidate_key)
            if score > best_score:
                best_value, best_score = candidate, score
        return Match(best_value, "Fuzzy", round(best_score, 3))

    def resolve(self, raw: Any) -> Match:
        text = clean_text(raw)
        if not text:
            return Match("", "Blank", 0.0)
        if text in self._cache:
            return self._cache[text]
        candidate = self.best_match(text)
        if candidate.method in {"Exact", "Normalized", "Alias"}:
            result = candidate
        elif candidate.value and candidate.score >= self.cutoff:
            result = candidate
        else:
            result = Match(text, "Unmatched", candidate.score)
        self._cache[text] = result
        return result


# -----------------------------------------------------------------------------
# Value parsers
# -----------------------------------------------------------------------------
EXCEL_EPOCH = datetime(1899, 12, 30)


def parse_date_value(raw: Any) -> tuple[date | None, str]:
    """Return (date, note). Note explains any assumption that was applied."""
    if raw is None or clean_text(raw) == "":
        return None, ""
    if isinstance(raw, datetime):
        return raw.date(), ""
    if isinstance(raw, date):
        return raw, ""
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        serial = float(raw)
        if 1 <= serial <= 80000:
            return (EXCEL_EPOCH + timedelta(days=serial)).date(), "Excel serial"
        return None, "Out-of-range numeric date"

    text = clean_text(raw)
    if "T" in text and re.match(r"^\d{4}-\d{2}-\d{2}T", text):
        text = text.split("T", 1)[0]
    text = text.replace(".", "-").replace("/", "-").replace(" ", "-")
    text = re.sub(r"-+", "-", text).strip("-")

    iso = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if iso:
        try:
            return date(int(iso.group(1)), int(iso.group(2)), int(iso.group(3))), ""
        except ValueError:
            return None, "Invalid ISO date"

    named = re.fullmatch(r"(\d{1,2})-([A-Za-z]+)-(\d{2,4})", text)
    if named and normalize(named.group(2))[:4] in {n[:4] for n in MONTH_NAMES}:
        month = next(
            (num for name, num in MONTH_NAMES.items() if normalize(named.group(2)) == name),
            None,
        )
        if month:
            year = int(named.group(3))
            year += 2000 if year < 100 else 0
            try:
                return date(year, month, int(named.group(1))), ""
            except ValueError:
                return None, "Invalid day for month"

    numeric = re.fullmatch(r"(\d{1,2})-(\d{1,2})-(\d{2,4})", text)
    if numeric:
        first, second = int(numeric.group(1)), int(numeric.group(2))
        year = int(numeric.group(3))
        year += 2000 if year < 100 else 0
        order = [(first, second), (second, first)]
        note = ""
        if not PREFER_DAY_FIRST_DATES:
            order.reverse()
        if order[0][1] > 12 and order[1][1] <= 12:
            order.reverse()
            note = "Read as month-first; day-first was impossible"
        day, month = order[0]
        try:
            return date(year, month, day), note
        except ValueError:
            return None, "Invalid calendar date"
    return None, "Unrecognised date format"


def parse_number(raw: Any) -> tuple[float | None, str]:
    if raw is None or clean_text(raw) == "":
        return None, ""
    if isinstance(raw, bool):
        return None, "Boolean is not a number"
    if isinstance(raw, (int, float)):
        return float(raw), ""
    text = clean_text(raw).replace(",", "").replace("%", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None, "No number found"
    return float(match.group()), "Extracted from text" if match.group() != text else ""


def tidy_number(value: float) -> Any:
    return int(value) if float(value).is_integer() else round(float(value), 4)


YES_TOKENS = {"yes", "y", "true", "1", "yes ", "available", "certified", "fluent", "done", "reviewed"}
NO_TOKENS = {"no", "n", "false", "0", "not available", "not certified", "not fluent", "nil", "none"}
NA_TOKENS = {"na", "n a", "not applicable", "nota"}


def parse_yes_no(raw: Any, allow_na: bool = False) -> tuple[str | None, str]:
    text = normalize(raw)
    if not text:
        return None, ""
    if text in YES_TOKENS:
        return "Yes", ""
    if text in NO_TOKENS:
        return "No", ""
    if allow_na and (text in NA_TOKENS or text.startswith("na ")):
        return "NA ( Incase of Automated Proctoring )", ""
    if text.startswith("yes"):
        return "Yes", ""
    if text.startswith("no") and not text.startswith("not "):
        return "No", ""
    return None, "Not a Yes/No value"


def parse_level(raw: Any) -> tuple[str | None, str]:
    if raw is None or clean_text(raw) == "":
        return None, ""
    value, _note = parse_number(raw)
    if value is None:
        return None, "Not a level number"
    text = f"{value:g}"
    if text not in {"1", "2", "2.5", "3", "3.5", "4", "4.5", "5", "5.5", "6", "6.5", "7", "8"}:
        return text, "Level outside the NSQF 1-8 list"
    return text, ""


URL_PATTERN = re.compile(r"https?://\S+", re.IGNORECASE)


def parse_link(raw: Any) -> tuple[str, str]:
    text = clean_text(raw)
    if not text:
        return "", ""
    if normalize(text) in NA_TOKENS | {"nil", "none", "not uploaded", "no link"}:
        return "", "Placeholder link removed"
    match = URL_PATTERN.search(text)
    if match:
        url = match.group().rstrip(").,;")
        return url, "" if url == text else "Extracted URL"
    if text.lower().startswith("www."):
        return "https://" + text, "Added scheme"
    return text, "Does not look like a URL"


# -----------------------------------------------------------------------------
# Master list loading
# -----------------------------------------------------------------------------
def _first_matching_column(headers: Sequence[Any], needles: Sequence[str]) -> int | None:
    normalized = [normalize(value) for value in headers]
    for needle in needles:
        target = normalize(needle)
        for index, header in enumerate(normalized):
            if header == target:
                return index
    for needle in needles:
        target = normalize(needle)
        for index, header in enumerate(normalized):
            if target and target in header:
                return index
    return None


def _iter_sheet_rows(path: Path, sheet_name: str | None = None):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [workbook[sheet_name]] if sheet_name else list(workbook.worksheets)
        for sheet in sheets:
            yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def load_single_column_list(path: Path, needles: Sequence[str], label: str) -> list[str]:
    if not path or not Path(path).is_file():
        logging.warning("%s master not found: %s", label, path)
        return []
    for _title, rows in _iter_sheet_rows(Path(path)):
        for header_index, row in enumerate(rows[:15]):
            column = _first_matching_column(row, needles)
            if column is None:
                continue
            values = [
                clean_text(item[column] if column < len(item) else None)
                for item in rows[header_index + 1:]
            ]
            values = [value for value in values if value]
            if values:
                logging.info("Loaded %d %s entries from %s", len(values), label, Path(path).name)
                return values
    logging.warning("Could not locate a %s column in %s", label, Path(path).name)
    return []


def load_state_districts(path: Path) -> dict[str, list[str]]:
    result: dict[str, list[str]] = defaultdict(list)
    if not path or not Path(path).is_file():
        logging.warning("State/District master not found: %s", path)
        return {}
    for _title, rows in _iter_sheet_rows(Path(path)):
        for header_index, row in enumerate(rows[:15]):
            state_col = _first_matching_column(row, ["State Name (In English)", "State Name", "State/ UT", "State"])
            district_col = _first_matching_column(row, ["District Name(In English)", "District Name", "District"])
            if state_col is None or district_col is None or state_col == district_col:
                continue
            for item in rows[header_index + 1:]:
                state = clean_text(item[state_col] if state_col < len(item) else None)
                district = clean_text(item[district_col] if district_col < len(item) else None)
                if state and district and district not in result[state]:
                    result[state].append(district)
            if result:
                logging.info(
                    "Loaded %d states / %d districts from %s",
                    len(result), sum(len(v) for v in result.values()), Path(path).name,
                )
                return dict(result)
    logging.warning("Could not locate State/District columns in %s", Path(path).name)
    return {}


def load_dcf_lookups(path: Path) -> dict[str, list[str]]:
    """Read the DCF LOOKUP sheet into {heading: [allowed values]}."""
    lookups: dict[str, list[str]] = {}
    if not path or not Path(path).is_file():
        logging.warning("DCF template not found: %s", path)
        return lookups
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if "LOOKUP" not in workbook.sheetnames:
            return lookups
        sheet = workbook["LOOKUP"]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        if not rows:
            return lookups
        width = max(len(row) for row in rows)
        for column in range(width):
            values = [
                clean_text(row[column] if column < len(row) else None) for row in rows
            ]
            values = [value for value in values if value]
            if len(values) < 2:
                continue
            heading, entries = values[0], values[1:]
            # Unheaded lookup columns (plain Yes/No lists) are keyed positionally.
            key = heading if normalize(heading) not in {"yes", "no", "accepted"} else f"column{column + 1}"
            lookups[key] = entries
        # The bare Yes/No/Accepted columns are still useful under stable keys.
        for column, key in ((1, "Yes/No"), (2, "Accepted/Rejected")):
            values = [clean_text(row[column] if column < len(row) else None) for row in rows]
            values = [value for value in values if value]
            if values:
                lookups.setdefault(key, values)
        return lookups
    finally:
        workbook.close()


# -----------------------------------------------------------------------------
# Column roles
# -----------------------------------------------------------------------------
TEXT_ROLE = "text"

COLUMN_ROLES: dict[str, str] = {
    normalize_header("Assessment Agency Name"): "agency",
    normalize_header("Assessment Agency (Folder)"): "agency",
    normalize_header("Batch ID"): "text",
    normalize_header("Is it a SIDH batch?"): "yesno",
    normalize_header("Is it a NSQF Aligned Batch?"): "nsqf_flag",
    normalize_header("Month of Batch Assessment"): "month",
    normalize_header("Year of Batch Assessment"): "year",
    normalize_header("Date on which Batch allocated by Awarding Body (DD-MM-YYYY)"): "date",
    normalize_header("Batch Accepted/Rejected"): "accepted",
    normalize_header("State/ UT"): "state",
    normalize_header("District"): "district",
    normalize_header("Assessment centre address"): "text",
    normalize_header("Funding Type"): "funding",
    normalize_header("Scheme Name"): "text",
    normalize_header("Training Type"): "training_type",
    normalize_header("Type of qualification"): "qualification_type",
    normalize_header("Sector"): "sector",
    normalize_header("Level"): "level",
    normalize_header("NQR code"): "text",
    normalize_header("Qualification name"): "text",
    normalize_header("Mode of assessment"): "mode",
    normalize_header("Primary Language of assessment"): "language",
    normalize_header("Name of Awarding Entity"): "text",
    normalize_header("Type of Awarding Entity"): "awarding_type",
    normalize_header("Start Date of Scheduled Assessment (DD-MM-YYYY)"): "date",
    normalize_header("Start Date of Actual Assessment (DD-MM-YYYY)"): "date",
    normalize_header("End Date of Actual Assessment (DD-MM-YYYY)"): "date",
    normalize_header("Assessed by valid ToA Certified Assessor"): "yesno",
    normalize_header("No. of Assessors Deployed for Assessment"): "number",
    normalize_header("Primary Assessor ID"): "text",
    normalize_header("Primary Assessor efficiency in language of assessment"): "yesno",
    normalize_header("Mode of Proctoring"): "proctoring",
    normalize_header("No. of candidates scheduled"): "number",
    normalize_header("No. of Candidates assessed"): "number",
    normalize_header("No. of Candidates passed"): "number",
    normalize_header("Average % Marks obtained"): "number",
    normalize_header("Has the result been analysed by Quality/Audit Team?"): "yesno_na",
    normalize_header("Date of result submission (DD-MM-YYYY)"): "date",
    normalize_header("Whether result is sent back by AB for correction"): "yesno",
    normalize_header("Assessment link"): "link",
    normalize_header("Has the videos been reviewed by Quality/Audit Team?"): "yesno_na",
    normalize_header("Remarks"): "text",
    # Traceability columns written by the combiner.
    normalize_header("Reporting Month (Folder)"): "skip",
    normalize_header("Source File Format Status"): "skip",
    normalize_header("Source File"): "skip",
    normalize_header("Source Sheet"): "skip",
    normalize_header("Source Row"): "skip",
}

ORIGINAL_TRACKED_ROLES = {"agency", "state", "district", "sector"}


def role_for_header(header: str) -> str:
    key = normalize_header(header)
    if key in COLUMN_ROLES:
        return COLUMN_ROLES[key]
    # Tolerate wording drift in the DCF headers.
    for known_key, role in COLUMN_ROLES.items():
        if known_key and (known_key in key or key in known_key) and len(key) > 6:
            return role
    text = normalize(header)
    if "date" in text:
        return "date"
    if text.startswith("no of") or "number of" in text:
        return "number"
    return TEXT_ROLE


# -----------------------------------------------------------------------------
# Cleaner
# -----------------------------------------------------------------------------
@dataclass
class Change:
    row: int
    agency: str
    column: str
    original: Any
    cleaned: Any
    change_type: str
    method: str
    score: float
    note: str = ""


@dataclass
class Cleaner:
    agencies: Vocabulary
    sectors: Vocabulary
    states: Vocabulary
    districts_by_state: dict[str, Vocabulary]
    all_districts: Vocabulary
    district_state: dict[str, str]
    lookups: dict[str, Vocabulary]
    month_hint: tuple[int, int] | None = None

    changes: list[Change] = field(default_factory=list)
    change_counts: Counter = field(default_factory=Counter)
    unmatched: Counter = field(default_factory=Counter)
    flags_per_row: list[str] = field(default_factory=list)
    truncated_log: bool = False

    def record(self, change: Change) -> None:
        self.change_counts[(change.column, change.change_type)] += 1
        if len(self.changes) < MAX_LOG_ROWS:
            self.changes.append(change)
        else:
            self.truncated_log = True

    # -- individual value cleaners -------------------------------------------
    def clean_vocab(self, raw: Any, vocab: Vocabulary, label: str) -> tuple[Any, str, str, float]:
        if not vocab:
            return clean_text(raw), "", "None", 0.0
        match = vocab.resolve(raw)
        if match.method == "Blank":
            return "", "", "Blank", 0.0
        if match.method == "Unmatched":
            self.unmatched[(label, clean_text(raw))] += 1
            return clean_text(raw), f"{label} not in master list", "Unmatched", match.score
        return match.value, "", match.method, match.score

    def clean_cell(
        self,
        role: str,
        raw: Any,
        row_context: dict[str, Any],
    ) -> tuple[Any, str, str, float, str]:
        """Return (value, flag, method, score, note)."""
        if role == "skip":
            return raw, "", "Skipped", 0.0, ""

        if role == "agency":
            value, flag, method, score = self.clean_vocab(raw, self.agencies, "Assessment Agency")
            return value, flag, method, score, ""

        if role == "sector":
            value, flag, method, score = self.clean_vocab(raw, self.sectors, "Sector")
            return value, flag, method, score, ""

        if role == "state":
            value, flag, method, score = self.clean_vocab(raw, self.states, "State/UT")
            row_context["state"] = value if method != "Unmatched" else ""
            return value, flag, method, score, ""

        if role == "district":
            return self.clean_district(raw, row_context)

        if role == "month":
            return self.clean_month(raw, row_context)

        if role == "year":
            return self.clean_year(raw, row_context)

        if role == "date":
            parsed, note = parse_date_value(raw)
            if parsed is None:
                if clean_text(raw) == "":
                    return "", "", "Blank", 0.0, ""
                self.unmatched[("Date", clean_text(raw))] += 1
                return clean_text(raw), f"Unparseable date ({note})", "Unmatched", 0.0, note
            same = isinstance(raw, (date, datetime)) and (
                raw.date() if isinstance(raw, datetime) else raw
            ) == parsed
            return parsed, "", "Exact" if same else "Parsed", 1.0, note

        if role == "number":
            value, note = parse_number(raw)
            if value is None:
                if clean_text(raw) == "":
                    return "", "", "Blank", 0.0, ""
                self.unmatched[("Number", clean_text(raw))] += 1
                return clean_text(raw), f"Not numeric ({note})", "Unmatched", 0.0, note
            tidy = tidy_number(value)
            same = isinstance(raw, (int, float)) and not isinstance(raw, bool) and float(raw) == value
            return tidy, "", "Exact" if same else "Parsed", 1.0, note

        if role in {"yesno", "yesno_na"}:
            value, note = parse_yes_no(raw, allow_na=(role == "yesno_na"))
            if value is None:
                if clean_text(raw) == "":
                    return "", "", "Blank", 0.0, ""
                self.unmatched[("Yes/No", clean_text(raw))] += 1
                return clean_text(raw), "Not a Yes/No value", "Unmatched", 0.0, note
            return value, "", "Exact" if clean_text(raw) == value else "Normalized", 1.0, ""

        if role == "nsqf_flag":
            vocab = self.lookups.get("nsqf")
            if vocab:
                text = normalize(raw)
                if text in {"yes", "y", "nsqf", "nsqf aligned"}:
                    return "NSQF Aligned", "", "Normalized", 1.0, ""
                if text in {"no", "n", "non nsqf", "not nsqf"}:
                    return "Non NSQF", "", "Normalized", 1.0, ""
                value, flag, method, score = self.clean_vocab(raw, vocab, "NSQF flag")
                return value, flag, method, score, ""
            return self.clean_cell("yesno", raw, row_context)

        if role == "accepted":
            vocab = self.lookups.get("accepted")
            if vocab:
                value, flag, method, score = self.clean_vocab(raw, vocab, "Accepted/Rejected")
                return value, flag, method, score, ""
            return clean_text(raw), "", "Trimmed", 0.0, ""

        if role == "level":
            value, note = parse_level(raw)
            if value is None:
                if clean_text(raw) == "":
                    return "", "", "Blank", 0.0, ""
                self.unmatched[("Level", clean_text(raw))] += 1
                return clean_text(raw), "Unrecognised NSQF level", "Unmatched", 0.0, note
            return value, note and f"{note}", "Parsed", 1.0, note

        if role == "link":
            value, note = parse_link(raw)
            flag = "Assessment link is not a URL" if note == "Does not look like a URL" else ""
            return value, flag, "Parsed" if note else "Trimmed", 1.0, note

        lookup_key = {
            "funding": "funding", "training_type": "training_type",
            "qualification_type": "qualification_type", "mode": "mode",
            "language": "language", "proctoring": "proctoring",
            "awarding_type": "awarding_type",
        }.get(role)
        if lookup_key:
            vocab = self.lookups.get(lookup_key)
            if vocab:
                label = lookup_key.replace("_", " ").title()
                value, flag, method, score = self.clean_vocab(raw, vocab, label)
                return value, flag, method, score, ""

        return clean_text(raw), "", "Trimmed", 0.0, ""

    def clean_district(self, raw: Any, row_context: dict[str, Any]) -> tuple[Any, str, str, float, str]:
        """Resolve the district inside its own state first, then nationally.

        A known State/UT narrows 784 districts down to a few dozen, so a match
        found inside the state is trusted at a lower similarity than a match
        found across the whole country.
        """
        text = clean_text(raw)
        if not text:
            return "", "", "Blank", 0.0, ""
        state = row_context.get("state") or ""
        state_vocab = self.districts_by_state.get(state) if state else None
        in_state = state_vocab.best_match(text) if state_vocab else Match("", "Fuzzy", 0.0)
        overall = self.all_districts.best_match(text) if self.all_districts else Match("", "Fuzzy", 0.0)

        relaxed_cutoff = max(0.60, FUZZY_CUTOFF_DISTRICT - 0.12)
        if in_state.value and in_state.score >= relaxed_cutoff:
            note = (
                ""
                if in_state.score >= FUZZY_CUTOFF_DISTRICT
                else f"Matched within {state} at reduced confidence"
            )
            return in_state.value, "", in_state.method, in_state.score, note

        if overall.value and overall.score >= FUZZY_CUTOFF_DISTRICT:
            inferred_state = self.district_state.get(normalize(overall.value), "")
            flag = note = ""
            if state and inferred_state and inferred_state != state:
                flag = f"District matches {inferred_state}, but State/UT says {state}"
            elif not state and inferred_state:
                row_context["inferred_state"] = inferred_state
                note = f"State inferred as {inferred_state}"
            return overall.value, flag, overall.method, overall.score, note

        self.unmatched[("District", text)] += 1
        return text, "District not in master list", "Unmatched", 0.0, ""

    def clean_month(self, raw: Any, row_context: dict[str, Any]) -> tuple[Any, str, str, float, str]:
        if isinstance(raw, (datetime, date)):
            value = raw.date() if isinstance(raw, datetime) else raw
            row_context["month"] = value.month
            row_context["year"] = value.year
            return MONTH_LABELS[value.month - 1], "", "Parsed", 1.0, "Read from a date value"
        text = normalize(raw)
        if not text:
            if self.month_hint:
                row_context["month"] = self.month_hint[1]
                return MONTH_LABELS[self.month_hint[1] - 1], "", "Filled", 1.0, "Filled from reporting month"
            folder = row_context.get("folder_period")
            if folder:
                row_context["month"] = folder[1]
                return MONTH_LABELS[folder[1] - 1], "", "Filled", 1.0, "Filled from source folder month"
            return "", "Month of batch assessment is blank", "Blank", 0.0, ""
        for name, number in MONTH_NAMES.items():
            if re.search(rf"\b{name}\b", text):
                row_context["month"] = number
                return MONTH_LABELS[number - 1], "", "Normalized", 1.0, ""
        number_value, _note = parse_number(raw)
        if number_value is not None and 1 <= int(number_value) <= 12:
            row_context["month"] = int(number_value)
            return MONTH_LABELS[int(number_value) - 1], "", "Parsed", 1.0, "Numeric month"
        self.unmatched[("Month", clean_text(raw))] += 1
        return clean_text(raw), "Unrecognised month", "Unmatched", 0.0, ""

    def clean_year(self, raw: Any, row_context: dict[str, Any]) -> tuple[Any, str, str, float, str]:
        if isinstance(raw, (datetime, date)):
            year = raw.year
            row_context["year"] = year
            return year, "", "Parsed", 1.0, "Read from a date value"
        value, _note = parse_number(raw)
        if value is None:
            fallback = row_context.get("year") or (self.month_hint[0] if self.month_hint else None)
            folder = row_context.get("folder_period")
            if fallback is None and folder:
                fallback = folder[0]
            if fallback:
                row_context["year"] = int(fallback)
                return int(fallback), "", "Filled", 1.0, "Filled from reporting month"
            if clean_text(raw) == "":
                return "", "Year of batch assessment is blank", "Blank", 0.0, ""
            self.unmatched[("Year", clean_text(raw))] += 1
            return clean_text(raw), "Unrecognised year", "Unmatched", 0.0, ""
        year = int(value)
        note = ""
        if year < 100:
            year += 2000
            note = "Two-digit year expanded"
        if not 2000 <= year <= 2100:
            self.unmatched[("Year", clean_text(raw))] += 1
            return year, "Year outside 2000-2100", "Unmatched", 0.0, note
        row_context["year"] = year
        return year, "", "Parsed" if note else "Exact", 1.0, note


def parse_period_text(value: Any) -> tuple[int, int] | None:
    text = clean_text(value)
    if not text:
        return None
    if isinstance(value, (datetime, date)):
        return (value.year, value.month)
    normalized = normalize(text)
    month = next((num for name, num in MONTH_NAMES.items() if re.search(rf"\b{name}\b", normalized)), None)
    years = re.findall(r"\b(20\d{2}|\d{2})\b", text)
    if month and years:
        year = int(years[-1])
        year += 2000 if year < 100 else 0
        return (year, month)
    numeric = re.fullmatch(r"(0?[1-9]|1[0-2])[-/](20\d{2}|\d{2})", text)
    if numeric:
        year = int(numeric.group(2))
        year += 2000 if year < 100 else 0
        return (year, int(numeric.group(1)))
    numeric = re.fullmatch(r"(20\d{2})[-/](0?[1-9]|1[0-2])", text)
    if numeric:
        return (int(numeric.group(1)), int(numeric.group(2)))
    return None


# -----------------------------------------------------------------------------
# Input reading
# -----------------------------------------------------------------------------
def select_data_sheet(workbook, preferred: str = ""):
    if preferred:
        if preferred not in workbook.sheetnames:
            raise ValueError(f"Input sheet not found: {preferred}")
        return workbook[preferred]
    for name in ("Master Data", "Data", "Combined Data"):
        if name in workbook.sheetnames:
            return workbook[name]
    candidates = [ws for ws in workbook.worksheets if ws.max_row > 1 and ws.max_column > 2]
    if not candidates:
        raise ValueError("No usable data sheet was found in the input workbook.")
    return max(candidates, key=lambda ws: ws.max_row * ws.max_column)


def find_header_row(rows: list[list[Any]]) -> int:
    best_row, best_score = 0, -1
    known = set(COLUMN_ROLES)
    for index, row in enumerate(rows[:30]):
        headers = {normalize_header(value) for value in row if clean_text(value)}
        score = len(headers & known) * 10 + len(headers)
        if score > best_score:
            best_row, best_score = index, score
    return best_row


# -----------------------------------------------------------------------------
# Output styling
# -----------------------------------------------------------------------------
def thin_border(color: str = "BFBFBF") -> Border:
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def style_header(ws, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = ws.cell(row, column)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border(NAVY)
    ws.row_dimensions[row].height = 40


def style_title(ws, title: str, columns: int, subtitle: str = "") -> int:
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(columns, 2))
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Aptos Display", size=14, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(columns, 2))
        cell = ws.cell(2, 1, subtitle)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(name="Aptos", size=9, italic=True, color=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        return 4
    return 3


# -----------------------------------------------------------------------------
# Main clean routine
# -----------------------------------------------------------------------------
def build_vocabularies() -> Cleaner:
    agency_names = load_single_column_list(
        AA_LIST_PATH, ["Assessment Agency Name", "Name of Assessment Agency", "Name"], "Assessment Agency"
    )
    sector_names = load_single_column_list(SECTORS_LIST_PATH, ["SECTOR", "Sector Name", "Sector"], "Sector")
    state_districts = load_state_districts(STATE_DISTRICT_PATH)
    lookups_raw = load_dcf_lookups(DCF_TEMPLATE_PATH)

    if not sector_names:
        sector_names = lookups_raw.get("Sector", [])

    states = sorted(state_districts) or []
    district_state: dict[str, str] = {}
    districts_by_state: dict[str, Vocabulary] = {}
    all_district_names: list[str] = []
    for state, districts in state_districts.items():
        districts_by_state[state] = Vocabulary(
            f"District ({state})", districts, cutoff=FUZZY_CUTOFF_DISTRICT
        )
        for district in districts:
            key = normalize(district)
            district_state.setdefault(key, state)
            if district not in all_district_names:
                all_district_names.append(district)

    def vocab_from_lookup(*keys: str, cutoff: float = FUZZY_CUTOFF) -> Vocabulary | None:
        for key in keys:
            for heading, values in lookups_raw.items():
                if normalize(heading) == normalize(key):
                    return Vocabulary(key, values, cutoff=cutoff)
        return None

    lookups: dict[str, Vocabulary] = {}
    for role, headings in {
        "funding": ("Funding",),
        "training_type": ("Training Type",),
        "qualification_type": ("Type Of Qualification",),
        "language": ("Language",),
        "mode": ("Mode of Assesment", "Mode of Assessment"),
        "proctoring": ("Mode or Proctoring", "Mode of Proctoring"),
        "accepted": ("Accepted/Rejected",),
    }.items():
        vocab = vocab_from_lookup(*headings)
        if vocab:
            lookups[role] = vocab
    lookups.setdefault("accepted", Vocabulary("Accepted/Rejected", ["Accepted", "Rejected"]))
    lookups.setdefault("awarding_type", Vocabulary("Type of Awarding Entity", ["NCVET Recognized", "Others"]))
    lookups.setdefault("nsqf", Vocabulary("NSQF flag", ["NSQF Aligned", "Non NSQF"]))

    sector_aliases = {}
    for value in lookups_raw.get("Sector", []):
        sector_aliases[value] = value  # resolved through the canonical list below

    cleaner = Cleaner(
        agencies=Vocabulary(
            "Assessment Agency", agency_names, aliases=AGENCY_NAME_ALIASES,
            cutoff=FUZZY_CUTOFF_AGENCY, normalizer=normalize_agency
        ),
        sectors=Vocabulary("Sector", sector_names, cutoff=FUZZY_CUTOFF),
        states=Vocabulary("State/UT", states, cutoff=FUZZY_CUTOFF),
        districts_by_state=districts_by_state,
        all_districts=Vocabulary("District", all_district_names, cutoff=FUZZY_CUTOFF_DISTRICT),
        district_state=district_state,
        lookups=lookups,
        month_hint=parse_period_text(REPORTING_MONTH_HINT),
    )
    return cleaner


def clean_workbook(
    input_path: Path,
    output_path: Path,
    preferred_sheet: str = "",
) -> dict[str, Any]:
    if not input_path.is_file():
        raise FileNotFoundError(f"INPUT_WORKBOOK is not a file: {input_path}")

    cleaner = build_vocabularies()

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        sheet = select_data_sheet(workbook, preferred_sheet)
        source_sheet_name = sheet.title
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not rows:
        raise ValueError("The input sheet is empty.")

    header_index = find_header_row(rows)
    headers = [clean_text(value) for value in rows[header_index]]
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise ValueError("No header row could be identified in the input sheet.")

    roles = [role_for_header(header) for header in headers]
    folder_month_col = next(
        (i for i, h in enumerate(headers) if normalize_header(h) == normalize_header("Reporting Month (Folder)")),
        None,
    )
    tracked = [
        (index, headers[index])
        for index, role in enumerate(roles)
        if role in ORIGINAL_TRACKED_ROLES
    ]

    output_headers = list(headers) + ["Cleaning Flags"]
    if KEEP_ORIGINAL_COLUMNS:
        output_headers += [f"Original {name}" for _index, name in tracked]

    cleaned_rows: list[list[Any]] = []
    date_columns: set[int] = set()
    rows_with_flags = 0

    for offset, raw_row in enumerate(rows[header_index + 1:], start=1):
        values = list(raw_row[: len(headers)]) + [None] * max(0, len(headers) - len(raw_row))
        if not any(clean_text(value) for value in values):
            continue
        excel_row = header_index + 1 + offset

        row_context: dict[str, Any] = {}
        if folder_month_col is not None and folder_month_col < len(values):
            row_context["folder_period"] = parse_period_text(values[folder_month_col])

        agency_display = ""
        agency_index = next((i for i, role in enumerate(roles) if role == "agency"), None)
        if agency_index is not None and agency_index < len(values):
            agency_display = clean_text(values[agency_index])

        originals = [values[index] for index, _name in tracked]
        cleaned_values: list[Any] = []
        flags: list[str] = []

        # State must be resolved before District, so process in header order but
        # give the state column a first pass when it appears after the district.
        order = list(range(len(headers)))
        state_positions = [i for i, role in enumerate(roles) if role == "state"]
        district_positions = [i for i, role in enumerate(roles) if role == "district"]
        if state_positions and district_positions and min(district_positions) < min(state_positions):
            order = state_positions + [i for i in order if i not in state_positions]

        results: dict[int, Any] = {}
        for index in order:
            role = roles[index]
            raw = values[index] if index < len(values) else None
            value, flag, method, score, note = cleaner.clean_cell(role, raw, row_context)
            if method == "Unmatched":
                value = "Missing"
            results[index] = value
            if role == "date" and isinstance(value, date):
                date_columns.add(index)
            original_text = display_value(raw)
            new_text = display_value(value)
            if method == "Skipped" or (not original_text and not new_text):
                pass
            elif method == "Unmatched":
                cleaner.record(
                    Change(
                        row=excel_row, agency=agency_display, column=headers[index],
                        original=original_text, cleaned=new_text,
                        change_type="Replaced with Missing",
                        method=method, score=score, note=note,
                    )
                )
            elif original_text != new_text:
                change_type = {"Filled": "Filled", "Blank": "Cleared"}.get(method, "Standardised")
                cleaner.record(
                    Change(
                        row=excel_row, agency=agency_display, column=headers[index],
                        original=original_text, cleaned=new_text,
                        change_type=change_type, method=method, score=score, note=note,
                    )
                )
            if flag:
                flags.append(f"{headers[index]}: {flag}")

        cleaned_values = [results.get(index, values[index]) for index in range(len(headers))]

        # Backfill a blank state when the district identified one.
        if "inferred_state" in row_context and state_positions:
            position = state_positions[0]
            if not clean_text(cleaned_values[position]):
                cleaned_values[position] = row_context["inferred_state"]
                cleaner.record(
                    Change(
                        row=excel_row, agency=agency_display, column=headers[position],
                        original="", cleaned=row_context["inferred_state"],
                        change_type="Filled", method="Inferred from district", score=1.0,
                        note="State derived from the matched district",
                    )
                )

        if flags:
            rows_with_flags += 1
        out_row = cleaned_values + ["; ".join(flags)]
        if KEEP_ORIGINAL_COLUMNS:
            out_row += [clean_text(value) for value in originals]
        cleaned_rows.append(out_row)

    if not cleaned_rows:
        raise ValueError("No populated data rows were found below the header row.")

    write_output(
        output_path=output_path,
        headers=output_headers,
        rows=cleaned_rows,
        date_columns=date_columns,
        cleaner=cleaner,
        input_path=input_path,
        source_sheet=source_sheet_name,
        rows_with_flags=rows_with_flags,
    )

    return {
        "input": str(input_path),
        "output": str(output_path),
        "source_sheet": source_sheet_name,
        "rows": len(cleaned_rows),
        "columns": len(output_headers),
        "changes": sum(cleaner.change_counts.values()),
        "rows_with_flags": rows_with_flags,
        "unmatched_values": len(cleaner.unmatched),
        "log_truncated": cleaner.truncated_log,
    }


def write_output(
    output_path: Path,
    headers: list[str],
    rows: list[list[Any]],
    date_columns: set[int],
    cleaner: Cleaner,
    input_path: Path,
    source_sheet: str,
    rows_with_flags: int,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "NCVET"
    workbook.properties.title = "Assessment Agency Cleaned Master"

    # --- Master Data ---------------------------------------------------------
    ws = workbook.create_sheet("Master Data")
    ws.sheet_view.showGridLines = False
    ws.append(headers)
    style_header(ws, 1, len(headers))
    flag_column = headers.index("Cleaning Flags") + 1
    for row in rows:
        ws.append(row)
    for index in date_columns:
        letter = get_column_letter(index + 1)
        for cell in ws[letter][1:]:
            cell.number_format = "DD-MM-YYYY"
    for row_number in range(2, ws.max_row + 1):
        cell = ws.cell(row_number, flag_column)
        if clean_text(cell.value):
            cell.fill = PatternFill("solid", fgColor=YELLOW)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    for index, header in enumerate(headers, start=1):
        text = normalize(header)
        width = 18
        if any(word in text for word in ("agency", "qualification", "address", "link", "scheme", "remarks", "flags", "source file")):
            width = 36
        ws.column_dimensions[get_column_letter(index)].width = width

    # --- Cleaning Log --------------------------------------------------------
    ws = workbook.create_sheet("Cleaning Log")
    subtitle = f"{len(cleaner.changes):,} change(s) listed"
    if cleaner.truncated_log:
        subtitle += f" (truncated at MAX_LOG_ROWS={MAX_LOG_ROWS:,}; see Change Summary for full totals)"
    header_row = style_title(ws, "Cleaning Log — every value the cleaner rewrote", 8, subtitle)
    log_headers = [
        "Master Data Row", "Assessment Agency", "Column", "Original Value",
        "Cleaned Value", "Change Type", "Match Method", "Confidence",
    ]
    for column, value in enumerate(log_headers, 1):
        ws.cell(header_row, column, value)
    style_header(ws, header_row, len(log_headers))
    for change in cleaner.changes:
        ws.append([
            change.row, change.agency, change.column, change.original,
            change.cleaned, change.change_type, change.method,
            round(change.score, 3) if change.score else "",
        ])
    for row_number in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row_number, 6).value
        if value == "Flagged - left as submitted":
            ws.cell(row_number, 6).fill = PatternFill("solid", fgColor=RED)
        elif value == "Standardised":
            ws.cell(row_number, 6).fill = PatternFill("solid", fgColor=GREEN)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:H{max(ws.max_row, header_row)}"
    for column, width in enumerate([16, 36, 34, 34, 34, 26, 16, 12], 1):
        ws.column_dimensions[get_column_letter(column)].width = width

    # --- Change Summary ------------------------------------------------------
    ws = workbook.create_sheet("Change Summary")
    header_row = style_title(ws, "Change Summary — totals by column and change type", 3)
    for column, value in enumerate(["Column", "Change Type", "Count"], 1):
        ws.cell(header_row, column, value)
    style_header(ws, header_row, 3)
    for (column_name, change_type), count in sorted(
        cleaner.change_counts.items(), key=lambda item: (-item[1], item[0])
    ):
        ws.append([column_name, change_type, count])
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:C{max(ws.max_row, header_row)}"
    for column, width in enumerate([44, 30, 12], 1):
        ws.column_dimensions[get_column_letter(column)].width = width

    # --- Unmatched Values ----------------------------------------------------
    ws = workbook.create_sheet("Unmatched Values")
    header_row = style_title(
        ws, "Unmatched Values — nothing in the master lists was close enough", 3,
        "These values were replaced with Missing in Master Data. Original submissions are retained here and in the Cleaning Log.",
    )
    for column, value in enumerate(["Field", "Submitted Value", "Occurrences"], 1):
        ws.cell(header_row, column, value)
    style_header(ws, header_row, 3)
    for (field_name, value), count in sorted(
        cleaner.unmatched.items(), key=lambda item: (item[0][0], -item[1])
    ):
        ws.append([field_name, value, count])
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:C{max(ws.max_row, header_row)}"
    for column, width in enumerate([24, 60, 14], 1):
        ws.column_dimensions[get_column_letter(column)].width = width

    # --- Run Log -------------------------------------------------------------
    ws = workbook.create_sheet("Run Log")
    header_row = style_title(ws, "Cleaning Run Log", 2, f"{APP_NAME} v{APP_VERSION}")
    for column, value in enumerate(["Item", "Value"], 1):
        ws.cell(header_row, column, value)
    style_header(ws, header_row, 2)
    entries = [
        ("Run at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Input workbook", str(input_path)),
        ("Input sheet", source_sheet),
        ("Output workbook", str(output_path)),
        ("Data rows cleaned", len(rows)),
        ("Rows carrying a cleaning flag", rows_with_flags),
        ("Total changes applied", sum(cleaner.change_counts.values())),
        ("Distinct unmatched values", len(cleaner.unmatched)),
        ("AA master", f"{AA_LIST_PATH} ({len(cleaner.agencies.values)} agencies)"),
        ("Sector master", f"{SECTORS_LIST_PATH} ({len(cleaner.sectors.values)} sectors)"),
        ("State/District master", f"{STATE_DISTRICT_PATH} ({len(cleaner.districts_by_state)} states, {len(cleaner.all_districts.values)} districts)"),
        ("DCF lookups", f"{DCF_TEMPLATE_PATH} ({len(cleaner.lookups)} controlled lists)"),
        ("Date convention", "Day-first (DD-MM-YYYY)" if PREFER_DAY_FIRST_DATES else "Month-first (MM-DD-YYYY)"),
        ("Fuzzy cutoffs", f"general {FUZZY_CUTOFF}, agency {FUZZY_CUTOFF_AGENCY}, district {FUZZY_CUTOFF_DISTRICT}"),
        ("Reporting month hint", REPORTING_MONTH_HINT or "(none)"),
    ]
    for item in entries:
        ws.append(list(item))
    for row_number in range(header_row + 1, ws.max_row + 1):
        for column in (1, 2):
            ws.cell(row_number, column).font = Font(name="Aptos", size=9)
            ws.cell(row_number, column).alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


# -----------------------------------------------------------------------------
# CLI
# -----------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Clean a combined AA workbook against the NCVET master lists."
    )
    parser.add_argument("--input", type=Path, default=INPUT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=OUTPUT_FILE)
    parser.add_argument("--sheet", default=INPUT_SHEET_NAME)
    parser.add_argument("--aa-list", type=Path, default=AA_LIST_PATH)
    parser.add_argument("--sectors", type=Path, default=SECTORS_LIST_PATH)
    parser.add_argument("--state-district", type=Path, default=STATE_DISTRICT_PATH)
    parser.add_argument("--dcf", type=Path, default=DCF_TEMPLATE_PATH)
    parser.add_argument("--month", default=REPORTING_MONTH_HINT)
    return parser


def apply_overrides(args: argparse.Namespace) -> None:
    global AA_LIST_PATH, SECTORS_LIST_PATH, STATE_DISTRICT_PATH
    global DCF_TEMPLATE_PATH, REPORTING_MONTH_HINT
    AA_LIST_PATH = Path(args.aa_list)
    SECTORS_LIST_PATH = Path(args.sectors)
    STATE_DISTRICT_PATH = Path(args.state_district)
    DCF_TEMPLATE_PATH = Path(args.dcf)
    REPORTING_MONTH_HINT = args.month


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s")
    apply_overrides(args)
    try:
        summary = clean_workbook(Path(args.input), Path(args.output), args.sheet)
    except Exception as exc:
        logging.exception("Cleaning failed")
        print(f"\nERROR: {exc}", file=sys.stderr)
        return 1
    print(f"\nDone. Cleaned {summary['rows']:,} rows from '{summary['source_sheet']}'.")
    print(f"Changes applied: {summary['changes']:,} | Rows flagged: {summary['rows_with_flags']:,}")
    print(f"Distinct unmatched values: {summary['unmatched_values']:,}")
    if summary["log_truncated"]:
        print("NOTE: the Cleaning Log was truncated; Change Summary holds the full totals.")
    print(f"Cleaned workbook: {summary['output']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
